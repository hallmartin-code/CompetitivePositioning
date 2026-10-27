"""Build the sample materials used by the tests and the sample run.

Cairnly is a fictional company. Every figure below is invented for testing.

    python tests/make_sample_deck.py            # writes into samples/
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

SAMPLE_SLIDES = [
    ("Cairnly", ["Supply ordering for independent dental practices", "Seed round  |  2026"], ""),
    ("The problem", [
        "Independent dental practices spend 6-8% of revenue on clinical supplies.",
        "Ordering is run by the office manager on top of the front desk job.",
        "Practices we interviewed reported 3 stockouts of a critical item per quarter.",
        "Office managers spend about 5 hours a week on ordering and reconciling invoices.",
    ], ""),
    ("How practices order today", [
        "1. Office manager walks the supply closet and updates a spreadsheet",
        "2. The distributor rep visits monthly and suggests a re-order",
        "3. Gaps are filled from Amazon Business or a second distributor at retail prices",
        "4. Invoices are matched to deliveries by hand at month end",
    ], "Speaker notes: none of the 22 practices we interviewed tracked usage per procedure. "
       "Most said the rep's suggested order is the default because nobody has time to check it."),
    ("Cairnly", [
        "Connects to the practice management system and reads scheduled procedures",
        "Forecasts supply usage per procedure and flags low stock before it runs out",
        "Places orders across distributors at the best contracted price",
    ], ""),
    ("Market", [
        "About 135,000 dental practices in the US; roughly 70% are independent (company estimate).",
        "Supply distribution is concentrated in two national distributors who run their own ordering portals.",
        "DSOs are consolidating practices and centralizing purchasing, which squeezes independents on price.",
    ], ""),
    ("Competition", [], ""),  # table added below
    ("Traction", [
        "14 practices in pilot across Ohio and Michigan",
        "5 paying customers at $249 per location per month",
        "Pilot practices cut supply spend 11% on average over 90 days (internal data, n=5)",
    ], ""),
    ("", [], ""),  # image-only product screenshot
    ("Team", [
        "Maya Okafor, CEO - 8 years running operations for a 30-location dental group",
        "Dev Raman, CTO - previously built ordering systems at a restaurant-supply marketplace",
    ], ""),
    ("The ask", ["Raising $1.5M seed to reach 150 paying locations in 18 months"], ""),
]

COMPETITION_TABLE = [
    ["Alternative", "How practices use it", "Gap"],
    ["Distributor portals", "Free with the distributor account", "Locked to one distributor's catalog and prices"],
    ["Spreadsheets", "Office manager tracks counts by hand", "No usage forecast; depends on one person"],
    ["Group purchasing orgs", "Negotiate discounts for members", "Price only; no ordering workflow"],
    ["Horizontal inventory apps", "Generic stock tracking", "No link to the procedure schedule"],
]

SUPPORT_PAGES = [
    "Customer interview notes - Cairnly pilot practices\n\n"
    "Interview 4, office manager, 3-chair practice, Columbus:\n"
    "\"I order whatever the rep puts in front of me. I don't have time to compare prices.\"\n\n"
    "Interview 9, practice owner, 5-chair practice, Ann Arbor:\n"
    "\"We ran out of composite twice this spring and had to reschedule patients.\"",
    "Interview 15, office manager, 2-location practice:\n"
    "\"Our spreadsheet lives on one computer. When I was on vacation nobody ordered anything.\"\n\n"
    "Pattern across 22 interviews: 17 of 22 practices buy from two or more distributors; "
    "none forecast usage from the schedule.",
]

THIN_SLIDES = [
    ("NovaCare AI", ["The future of healthcare"]),
    ("Vision", ["We are building an AI platform for healthcare.", "Huge opportunity."]),
    ("Team", ["Experienced founders with a passion for innovation."]),
    ("Ask", ["Raising $2M."]),
]


def _screenshot_png() -> bytes:
    img = Image.new("RGB", (1200, 700), "#F7F7F7")
    d = ImageDraw.Draw(img)
    d.rectangle([40, 40, 1160, 660], outline="#CBD6E2", width=4)
    for i in range(6):
        d.rectangle([80, 100 + i * 90, 1120, 160 + i * 90], fill="#FFFFFF", outline="#CBD6E2")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def build_sample_pptx(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, lines, notes in SAMPLE_SLIDES:
        if not title and not lines:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(io.BytesIO(_screenshot_png()), Inches(1), Inches(0.8), width=Inches(11))
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = title
        if title == "Competition":
            rows, cols = len(COMPETITION_TABLE), len(COMPETITION_TABLE[0])
            table = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.5), Inches(4)).table
            for r, row in enumerate(COMPETITION_TABLE):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        if lines:
            box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5))
            tf = box.text_frame
            tf.word_wrap = True
            for i, line in enumerate(lines):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = line
                para.font.size = Pt(22)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def build_sample_pdf(path: Path) -> Path:
    """The same deck as a text PDF, as a founder exporting from Keynote would send."""
    doc = fitz.open()
    for title, lines, notes in SAMPLE_SLIDES:
        page = doc.new_page(width=960, height=540)
        if not title and not lines:
            page.insert_image(fitz.Rect(60, 40, 900, 500), stream=_screenshot_png())
            continue
        page.insert_text((60, 80), title, fontsize=34)
        body = lines if title != "Competition" else [" | ".join(r) for r in COMPETITION_TABLE]
        page.insert_textbox(fitz.Rect(60, 120, 900, 520), "\n".join(body), fontsize=16)
    doc.save(path)
    return path


def build_support_pdf(path: Path) -> Path:
    doc = fitz.open()
    for text in SUPPORT_PAGES:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(60, 60, 550, 760), text, fontsize=11)
    doc.save(path)
    return path


def build_thin_pptx(path: Path) -> Path:
    prs = Presentation()
    for title, lines in THIN_SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(lines)
    prs.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "samples"
    out.mkdir(parents=True, exist_ok=True)
    for p in (build_sample_pptx(out / "Cairnly_Seed_Deck.pptx"),
              build_sample_pdf(out / "Cairnly_Seed_Deck.pdf"),
              build_support_pdf(out / "Cairnly_customer_interviews.pdf"),
              build_thin_pptx(out / "NovaCare_thin_deck.pptx")):
        print(p)
